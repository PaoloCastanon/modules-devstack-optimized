#!/usr/bin/env python3
"""Apply evidence-backed Sprint 2 status to a copy of the planning book."""
from pathlib import Path
import subprocess
import sys
from openpyxl import load_workbook
from openpyxl.workbook.properties import CalcProperties

root = Path(__file__).resolve().parents[1]
source = root / 'Planificacion_Trabajo_Titulo_OpenStack_Placement_Glance.xlsx'
backup = root / 'Planificacion_Trabajo_Titulo_OpenStack_Placement_Glance.pre-sprint2.xlsx'
target = root / 'Planificacion_Trabajo_Titulo_OpenStack_Placement_Glance_Sprint2.xlsx'
assert source.exists() and backup.exists()
assert source.read_bytes() == backup.read_bytes(), 'Original/backup differ'
subprocess.run([sys.executable, str(root / 'scripts/verify_sprint2.py')], check=True)
book = load_workbook(source)
assert book.sheetnames == ['Resumen','Objetivos_SMART','Requerimientos','Backlog','Trazabilidad','Sprints','Dashboard','Listas']
backlog = book['Backlog']
data = {
    'B01': ('Hecho','evidence/sprint2/environment/vm-audit.txt; infra/lab-vm/','VM Ubuntu 24.04.5 operativa: 4 vCPU, 8 GiB, disco 55 GiB; Terraform y Ansible.'),
    'B02': ('Hecho','evidence/sprint2/devstack/service-list.json; infra/devstack/versions.lock','DevStack instalado; Keystone, Nova, Placement, Glance y Neutron verificados; Ansible repetido.'),
    'B03': ('Hecho','evidence/sprint2/tests/boot-states.txt; evidence/sprint2/tests/smoke-verify.txt','Smoke completo; instancia CirrOS BUILD→ACTIVE y limpieza.'),
    'B04': ('Hecho','docs/sprint2/placement-architecture.md; evidence/sprint2/placement/nova-requests.log','Arquitectura contrastada con API, MySQL y tráfico real Nova→Placement.'),
    'B05': ('Hecho','docs/sprint2/placement-api-scope.md; evidence/sprint2/placement/versions.response.json','Inventario dentro/fuera; servidor 1.0–1.39; baseline y microversiones Nova 1.0/1.28/1.36.'),
    'B06': ('Hecho','docs/sprint2/nova-placement-flow.md; evidence/sprint2/placement/nova-requests.log','GET candidatos 1.36, GET allocations 1.28, PUT 1.36, DELETE 1.0; boot ACTIVE.'),
    'B19': ('Hecho','docs/sprint2/glance-api-scope.md; evidence/sprint2/glance/backend-audit.txt','Image API v2.18, auth 401, store real, upload/download/patch/delete y boot desde Glance.'),
}
changed = set()
for row in range(4, backlog.max_row + 1):
    key = backlog.cell(row, 1).value
    if key in data:
        state, evidence, observation = data[key]
        assert backlog.cell(row, 11).value == 'S2'
        backlog.cell(row, 12, state)
        backlog.cell(row, 16, evidence)
        backlog.cell(row, 17, observation)
        changed.add(key)
assert changed == set(data), changed
objectives = book['Objetivos_SMART']
assert objectives['A4'].value == 'OE1' and objectives['A5'].value == 'OE2'
objectives['K4'] = 'Cumplido'
objectives['L4'] = 'Arquitectura, APIs, dependencias y flujos críticos documentados con baseline y traza real; ver docs/sprint2/sprint2-evidence.md.'
objectives['K5'] = 'Cumplido'
objectives['L5'] = 'VM de un nodo reproducible, DevStack verificado, smoke completo y baseline original capturado.'
sprints = book['Sprints']
assert sprints['A5'].value == 'S2'
sprints['J5'] = 'Completado'
sprints['K5'] = '7/7 elementos S2 y 33/33 story points aceptados con evidencia; feedback formal del profesor no disponible.'
book.calculation = CalcProperties(fullCalcOnLoad=True)
book.save(target)
check = load_workbook(target)
assert check.sheetnames == book.sheetnames
assert sum(len(s._charts) for s in check) == 5
assert sum(c.data_type == 'f' for s in check for row in s for c in row) == sum(c.data_type == 'f' for s in book for row in s for c in row)
assert sum(len(s.data_validations.dataValidation) for s in check) == sum(len(s.data_validations.dataValidation) for s in book)
print(f'Updated {target.name}: {", ".join(sorted(changed))}; formulas/charts/validations preserved')
